from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "xTB4_one_page_performance_summary.docx"
BLUE = RGBColor(31, 78, 120)
DARK = RGBColor(11, 37, 69)
GRAY = RGBColor(89, 89, 89)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def set_cell_margins(cell, top=70, start=110, bottom=70, end=110):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); mar = tc_pr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tc_pr.append(mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def set_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr(); tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW'); tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width)); tc_w.set(qn('w:type'), 'dxa')

def set_run(run, size=9.2, bold=False, color=None):
    run.font.name = 'Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = color

def paragraph(doc_or_cell, text='', size=9.2, bold=False, color=None, after=3, before=0, align=None):
    p = doc_or_cell.add_paragraph() if hasattr(doc_or_cell, 'add_paragraph') else doc_or_cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.05
    if align is not None: p.alignment = align
    r = p.add_run(text); set_run(r, size, bold, color)
    return p

def add_heading(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper()); set_run(r, 9.5, True, BLUE)

def add_bullet(doc, lead, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.16); p.paragraph_format.first_line_indent = Inches(-.12)
    p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.line_spacing = 1.03
    r = p.add_run('• '); set_run(r, 9.1, True, DARK)
    r = p.add_run(lead); set_run(r, 9.1, True, DARK)
    r = p.add_run(text); set_run(r, 9.1, False, None)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.62); section.bottom_margin = Inches(.58)
section.left_margin = Inches(.72); section.right_margin = Inches(.72)
section.header_distance = Inches(.3); section.footer_distance = Inches(.3)

styles = doc.styles
styles['Normal'].font.name = 'Calibri'; styles['Normal'].font.size = Pt(9.2)

p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
r = p.add_run('xTB4 BAND-GAP MODEL — PERFORMANCE SUMMARY'); set_run(r, 18, True, DARK)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
r = p.add_run('DFT-free fragment features replace original DFT HOMO–LUMO inputs'); set_run(r, 10.2, False, GRAY)

# Key takeaway callout
table = doc.add_table(rows=1, cols=1); table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = table.cell(0, 0); set_cell_shading(cell, 'EAF1F8'); set_cell_margins(cell, 95, 130, 95, 130)
p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.03
r = p.add_run('Key finding. '); set_run(r, 10, True, DARK)
r = p.add_run('On the matched 216-COF complete-xTB cohort, standalone xTB4 TabPFN reaches R² = 0.624 and MAE = 0.381 eV—essentially matching the DFT-fragment baseline R² (0.625) while reducing MAE by 0.029 eV.'); set_run(r, 10, False, DARK)

add_heading(doc, 'Performance on identical 216-COF / 72-family grouped folds')
tbl = doc.add_table(rows=1, cols=4); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.style = 'Table Grid'
widths = [4900, 2190, 1050, 1220]
for i, label in enumerate(['Model / input representation', 'What changed', 'R²', 'MAE (eV)']):
    c = tbl.rows[0].cells[i]; set_width(c, widths[i]); set_cell_shading(c, 'DCE6F1'); set_cell_margins(c)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(label), 8.4, True, DARK)
rows = [
    ('DFT baseline — Random Forest', 'Original DFT fragment gaps + structure', '0.625', '0.410'),
    ('xTB3 — Random Forest', 'Raw xTB gaps + structure', '0.591', '0.433'),
    ('xTB4 — TabPFN v2', '92 rich xTB + structure features', '0.624', '0.381'),
    ('xTB4 — 75% TabPFN / 25% RF*', 'Same xTB4 features; blended predictions', '0.640*', '0.380*'),
]
for values in rows:
    cells = tbl.add_row().cells
    for i, value in enumerate(values):
        set_width(cells[i], widths[i]); set_cell_margins(cells[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i >= 2 else WD_ALIGN_PARAGRAPH.LEFT
        set_run(p.add_run(value), 8.5, i == 2 and values[0].startswith('xTB4 — TabPFN'), DARK if i == 0 else None)

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
r = p.add_run('* Exploratory result: blend weights were chosen after reviewing these folds; confirm with a locked test set or nested cross-validation.'); set_run(r, 7.8, False, GRAY)

add_heading(doc, 'What xTB4 changes')
add_bullet(doc, 'Input replacement: ', 'removes the original DFT fragment HOMO–LUMO/gap inputs and uses GFN2-xTB calculations on the same node, linker, and node+linker fragments.')
add_bullet(doc, 'Richer representation: ', 'adds HOMO/LUMO, IP/EA, electronegativity, hardness, dipoles, charge and bond-order summaries, orbital alignment, composition, and lattice-geometry descriptors.')
add_bullet(doc, 'Model change: ', 'TabPFN v2 is used for the 92 usable numerical features. It improves only when paired with the richer xTB4 representation; TabPFN with the older xTB3 features scored R² = 0.578.')

add_heading(doc, 'Data, validation, and interpretation')
add_bullet(doc, 'Cohort: ', '315 COFs have labels in bandgap_CoRE_COF_v5.0_DFTfree_features.csv. The benchmark uses 216 because all three required rich-xTB fragment calculations succeeded; 99 were excluded for missing fragment geometries, timeouts, or xTB convergence failures—not for their band gaps.')
add_bullet(doc, 'Validation: ', 'five-fold family-grouped cross-validation. In each fold, 172–173 COFs train the model and 43–44 held-out COFs are predicted. Family and row overlap are both asserted to be zero.')
add_bullet(doc, 'Meaning: ', 'R² measures explained variation (higher is better); MAE is mean absolute band-gap error (lower is better). The 0.381 eV MAE means a typical held-out prediction differs from the reference by about 0.38 eV.')

add_heading(doc, 'Current conclusion and next step')
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.04
r = p.add_run('Conclusion: '); set_run(r, 9.2, True, DARK)
r = p.add_run('xTB4 demonstrates that a DFT-free electronic-feature pipeline can match the DFT-fragment baseline on the complete-xTB cohort. '); set_run(r, 9.2)
r = p.add_run('Next step: '); set_run(r, 9.2, True, DARK)
r = p.add_run('recover or provide a validated fallback for the 99 incomplete xTB records, then repeat the same family-grouped evaluation on all 315 labeled COFs.'); set_run(r, 9.2)

footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_before = Pt(0); footer.paragraph_format.space_after = Pt(0)
set_run(footer.add_run('xTB4 performance brief | August 2026'), 7.5, False, GRAY)
doc.save(OUT)
print(OUT)
